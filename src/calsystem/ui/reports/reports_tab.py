"""
Reports tab - for generating PDF reports and Excel exports.
"""

from datetime import datetime
from typing import Optional, List, Dict, Any

from PyQt6.QtWidgets import (
    QWidget,
    QVBoxLayout,
    QHBoxLayout,
    QGroupBox,
    QPushButton,
    QTableWidget,
    QTableWidgetItem,
    QHeaderView,
    QLabel,
    QLineEdit,
    QSplitter,
    QComboBox,
    QTextEdit,
    QDateEdit,
    QFormLayout,
    QFileDialog,
    QMessageBox,
)
from PyQt6.QtCore import Qt, QDate
from PyQt6.QtGui import QColor
from loguru import logger

from calsystem.database.connection import get_db
from calsystem.database.models import CalibrationSession, DUT, Procedure, TestResult, TestPoint


class ReportsTab(QWidget):
    """Tab for generating reports and exporting data."""

    def __init__(self):
        super().__init__()
        self._current_session_id: Optional[int] = None
        self._current_session_data: Optional[Dict] = None
        self._init_ui()

    def showEvent(self, event):
        """Called when tab becomes visible."""
        super().showEvent(event)
        self._on_search()  # Refresh results

    def _init_ui(self):
        """Initialize the UI."""
        layout = QVBoxLayout(self)

        # Search/filter bar
        filter_group = QGroupBox("Search Calibration Records")
        filter_layout = QHBoxLayout(filter_group)

        filter_layout.addWidget(QLabel("Asset Number:"))
        self.asset_filter = QLineEdit()
        self.asset_filter.setPlaceholderText("Enter asset number...")
        filter_layout.addWidget(self.asset_filter)

        filter_layout.addWidget(QLabel("Work Order:"))
        self.workorder_filter = QLineEdit()
        filter_layout.addWidget(self.workorder_filter)

        filter_layout.addWidget(QLabel("Date Range:"))
        self.date_from = QDateEdit()
        self.date_from.setCalendarPopup(True)
        self.date_from.setDate(QDate.currentDate().addMonths(-1))
        filter_layout.addWidget(self.date_from)

        filter_layout.addWidget(QLabel("to"))
        self.date_to = QDateEdit()
        self.date_to.setCalendarPopup(True)
        self.date_to.setDate(QDate.currentDate())
        filter_layout.addWidget(self.date_to)

        self.search_btn = QPushButton("Search")
        self.search_btn.clicked.connect(self._on_search)
        filter_layout.addWidget(self.search_btn)

        layout.addWidget(filter_group)

        # Main splitter
        splitter = QSplitter(Qt.Orientation.Horizontal)

        # Left - Session list
        list_widget = QWidget()
        list_layout = QVBoxLayout(list_widget)
        list_layout.setContentsMargins(0, 0, 0, 0)

        list_group = QGroupBox("Calibration Sessions")
        list_inner_layout = QVBoxLayout(list_group)

        self.sessions_table = QTableWidget()
        self.sessions_table.setColumnCount(5)
        self.sessions_table.setHorizontalHeaderLabels(
            ["Date", "Asset Number", "Work Order", "Technician", "Result"]
        )
        self.sessions_table.horizontalHeader().setSectionResizeMode(
            QHeaderView.ResizeMode.Stretch
        )
        self.sessions_table.setSelectionBehavior(
            QTableWidget.SelectionBehavior.SelectRows
        )
        self.sessions_table.itemSelectionChanged.connect(self._on_session_selected)
        list_inner_layout.addWidget(self.sessions_table)

        list_layout.addWidget(list_group)
        splitter.addWidget(list_widget)

        # Right - Session details and report options
        details_widget = QWidget()
        details_layout = QVBoxLayout(details_widget)
        details_layout.setContentsMargins(0, 0, 0, 0)

        # Session details
        details_group = QGroupBox("Session Details")
        details_inner_layout = QFormLayout(details_group)

        self.detail_asset = QLabel("--")
        details_inner_layout.addRow("Asset Number:", self.detail_asset)

        self.detail_dut = QLabel("--")
        details_inner_layout.addRow("DUT:", self.detail_dut)

        self.detail_workorder = QLabel("--")
        details_inner_layout.addRow("Work Order:", self.detail_workorder)

        self.detail_date = QLabel("--")
        details_inner_layout.addRow("Date:", self.detail_date)

        self.detail_technician = QLabel("--")
        details_inner_layout.addRow("Technician:", self.detail_technician)

        self.detail_result = QLabel("--")
        details_inner_layout.addRow("Result:", self.detail_result)

        self.detail_procedure = QLabel("--")
        details_inner_layout.addRow("Procedure:", self.detail_procedure)

        details_layout.addWidget(details_group)

        # Test results summary
        results_group = QGroupBox("Test Results")
        results_layout = QVBoxLayout(results_group)

        self.results_table = QTableWidget()
        self.results_table.setColumnCount(6)
        self.results_table.setHorizontalHeaderLabels(
            ["Test Point", "Nominal", "Measured", "Deviation", "Tolerance", "Status"]
        )
        self.results_table.horizontalHeader().setSectionResizeMode(
            QHeaderView.ResizeMode.Stretch
        )
        results_layout.addWidget(self.results_table)

        details_layout.addWidget(results_group)

        # Report generation
        report_group = QGroupBox("Generate Report")
        report_layout = QVBoxLayout(report_group)

        template_layout = QHBoxLayout()
        template_layout.addWidget(QLabel("Template:"))
        self.template_combo = QComboBox()
        self.template_combo.addItems(["Standard Report", "Detailed Report", "Summary Only"])
        template_layout.addWidget(self.template_combo)
        template_layout.addStretch()
        report_layout.addLayout(template_layout)

        btn_layout = QHBoxLayout()

        self.pdf_btn = QPushButton("Generate PDF")
        self.pdf_btn.clicked.connect(self._on_generate_pdf)
        btn_layout.addWidget(self.pdf_btn)

        self.excel_btn = QPushButton("Export to Excel")
        self.excel_btn.clicked.connect(self._on_export_excel)
        btn_layout.addWidget(self.excel_btn)

        self.print_btn = QPushButton("Print")
        self.print_btn.clicked.connect(self._on_print)
        btn_layout.addWidget(self.print_btn)

        btn_layout.addStretch()
        report_layout.addLayout(btn_layout)

        details_layout.addWidget(report_group)

        splitter.addWidget(details_widget)

        # Set splitter sizes
        splitter.setSizes([400, 500])

        layout.addWidget(splitter)

    def _on_search(self):
        """Search for calibration sessions."""
        asset = self.asset_filter.text().strip()
        workorder = self.workorder_filter.text().strip()
        date_from = self.date_from.date().toPyDate()
        date_to = self.date_to.date().toPyDate()

        logger.info(
            f"Searching sessions: asset={asset}, wo={workorder}, "
            f"from={date_from}, to={date_to}"
        )

        self.sessions_table.setRowCount(0)

        db = get_db()
        if not db.is_connected:
            return

        try:
            with db.session() as session:
                query = session.query(CalibrationSession).join(DUT)

                # Apply filters
                if asset:
                    query = query.filter(DUT.asset_number.ilike(f"%{asset}%"))

                if workorder:
                    query = query.filter(CalibrationSession.work_order.ilike(f"%{workorder}%"))

                # Date range
                from datetime import datetime, timedelta
                date_from_dt = datetime.combine(date_from, datetime.min.time())
                date_to_dt = datetime.combine(date_to, datetime.max.time())
                query = query.filter(CalibrationSession.started_at >= date_from_dt)
                query = query.filter(CalibrationSession.started_at <= date_to_dt)

                # Order by date descending
                results = query.order_by(CalibrationSession.started_at.desc()).limit(100).all()

                for cal_session in results:
                    row = self.sessions_table.rowCount()
                    self.sessions_table.insertRow(row)

                    # Date
                    date_str = cal_session.started_at.strftime("%Y-%m-%d %H:%M") if cal_session.started_at else "--"
                    date_item = QTableWidgetItem(date_str)
                    date_item.setData(Qt.ItemDataRole.UserRole, cal_session.id)
                    self.sessions_table.setItem(row, 0, date_item)

                    # Asset
                    self.sessions_table.setItem(row, 1, QTableWidgetItem(cal_session.dut.asset_number if cal_session.dut else "--"))

                    # Work Order
                    self.sessions_table.setItem(row, 2, QTableWidgetItem(cal_session.work_order or "--"))

                    # Technician
                    self.sessions_table.setItem(row, 3, QTableWidgetItem(cal_session.technician_name or "--"))

                    # Result - handle status being either enum or string
                    status_val = cal_session.status.value if hasattr(cal_session.status, 'value') else str(cal_session.status or "")
                    result = (cal_session.overall_result or status_val).title()
                    result_item = QTableWidgetItem(result)
                    if result.lower() == "pass" or result.lower() == "completed":
                        result_item.setForeground(Qt.GlobalColor.darkGreen)
                    elif result.lower() == "fail":
                        result_item.setForeground(Qt.GlobalColor.red)
                    elif result.lower() == "aborted":
                        result_item.setForeground(Qt.GlobalColor.darkYellow)
                    self.sessions_table.setItem(row, 4, result_item)

                logger.debug(f"Found {len(results)} sessions")

        except Exception as e:
            logger.error(f"Search failed: {e}")

    def _on_session_selected(self):
        """Handle session selection."""
        selected = self.sessions_table.selectedItems()
        if not selected:
            return

        row = selected[0].row()
        date_item = self.sessions_table.item(row, 0)
        session_id = date_item.data(Qt.ItemDataRole.UserRole)

        if not session_id:
            return

        self._current_session_id = session_id
        self._load_session_details(session_id)

    def _load_session_details(self, session_id: int):
        """Load session details from database."""
        db = get_db()
        if not db.is_connected:
            return

        try:
            with db.session() as session:
                cal_session = session.query(CalibrationSession).filter(
                    CalibrationSession.id == session_id
                ).first()

                if not cal_session:
                    return

                # Update details panel
                dut = cal_session.dut
                procedure = cal_session.procedure

                self.detail_asset.setText(dut.asset_number if dut else "--")
                self.detail_dut.setText(f"{dut.make} {dut.model}" if dut else "--")
                self.detail_workorder.setText(cal_session.work_order or "--")
                self.detail_date.setText(
                    cal_session.started_at.strftime("%Y-%m-%d %H:%M") if cal_session.started_at else "--"
                )
                self.detail_technician.setText(cal_session.technician_name or "--")

                status_val = cal_session.status.value if hasattr(cal_session.status, 'value') else str(cal_session.status or "")
                result = (cal_session.overall_result or status_val).title()
                self.detail_result.setText(result)
                if result.lower() == "pass":
                    self.detail_result.setStyleSheet("color: green; font-weight: bold;")
                elif result.lower() == "fail":
                    self.detail_result.setStyleSheet("color: red; font-weight: bold;")
                else:
                    self.detail_result.setStyleSheet("")

                self.detail_procedure.setText(procedure.name if procedure else "--")

                # Store session data for export
                self._current_session_data = {
                    "id": session_id,
                    "asset": dut.asset_number if dut else "",
                    "dut_info": f"{dut.make} {dut.model} S/N: {dut.serial_number}" if dut else "",
                    "work_order": cal_session.work_order,
                    "date": cal_session.started_at,
                    "technician": cal_session.technician_name,
                    "result": result,
                    "procedure": procedure.name if procedure else "",
                }

                # Load test results
                self._load_test_results(session_id)

                logger.debug(f"Loaded session {session_id}")

        except Exception as e:
            logger.error(f"Failed to load session details: {e}")

    def _load_test_results(self, session_id: int):
        """Load test results for a session."""
        self.results_table.setRowCount(0)

        db = get_db()
        if not db.is_connected:
            return

        try:
            with db.session() as session:
                # Use raw SQL to avoid enum mapping issues
                from sqlalchemy import text
                query = text("""
                    SELECT
                        tr.id, tr.measured_value, tr.status,
                        tp.id as tp_id, tp.description, tp.nominal_value, tp.unit,
                        tp.tol_pct_reading, tp.tol_pct_range, tp.tol_pct_span,
                        tp.tol_digits, tp.tol_absolute, tp.tol_resolution,
                        tp.tol_range_value, tp.tol_span_value,
                        tp.tolerance_value, tp.tolerance_type
                    FROM test_results tr
                    JOIN test_points tp ON tr.test_point_id = tp.id
                    WHERE tr.session_id = :session_id
                    ORDER BY tr.executed_at
                """)
                results = session.execute(query, {"session_id": session_id}).fetchall()

                from calsystem.utils.tolerance import ToleranceSpec

                for result in results:
                    row = self.results_table.rowCount()
                    self.results_table.insertRow(row)

                    # Test Point
                    test_name = result.description or f"Point {row + 1}"
                    self.results_table.setItem(row, 0, QTableWidgetItem(test_name))

                    # Nominal
                    nominal_str = f"{result.nominal_value} {result.unit}" if result.nominal_value else "--"
                    self.results_table.setItem(row, 1, QTableWidgetItem(nominal_str))

                    # Measured
                    measured_str = f"{result.measured_value:.6g}" if result.measured_value else "--"
                    self.results_table.setItem(row, 2, QTableWidgetItem(measured_str))

                    # Deviation
                    if result.nominal_value is not None and result.measured_value is not None:
                        deviation = result.measured_value - result.nominal_value
                        if result.nominal_value != 0:
                            dev_pct = (deviation / result.nominal_value) * 100
                            dev_str = f"{dev_pct:+.4f}%"
                        else:
                            dev_str = f"{deviation:+.6g}"
                    else:
                        dev_str = "--"
                    self.results_table.setItem(row, 3, QTableWidgetItem(dev_str))

                    # Tolerance
                    spec = ToleranceSpec(
                        pct_reading=result.tol_pct_reading or 0,
                        pct_range=result.tol_pct_range or 0,
                        pct_span=result.tol_pct_span or 0,
                        digits=result.tol_digits or 0,
                        absolute=result.tol_absolute or 0,
                        resolution=result.tol_resolution or 0,
                        range_value=result.tol_range_value or 0,
                        span_value=result.tol_span_value or 0,
                    )
                    if spec.is_empty() and result.tolerance_value:
                        legacy_type = result.tolerance_type or "percent"
                        spec = ToleranceSpec.from_legacy(result.tolerance_value, legacy_type)

                    if not spec.is_empty():
                        calculated = spec.calculate(result.nominal_value or 0)
                        tol_str = f"{spec.format_spec()} (±{calculated:g})"
                    else:
                        tol_str = "--"
                    self.results_table.setItem(row, 4, QTableWidgetItem(tol_str))

                    # Status (already a string from raw query)
                    status = str(result.status).title() if result.status else "--"
                    status_item = QTableWidgetItem(status)
                    if status.lower() == "pass":
                        status_item.setForeground(Qt.GlobalColor.darkGreen)
                    elif status.lower() == "fail":
                        status_item.setForeground(Qt.GlobalColor.red)
                    self.results_table.setItem(row, 5, status_item)

                logger.debug(f"Loaded {len(results)} test results")

        except Exception as e:
            logger.error(f"Failed to load test results: {e}")

    def _on_generate_pdf(self):
        """Generate PDF report."""
        if not self._current_session_id or not self._current_session_data:
            QMessageBox.warning(
                self, "No Selection", "Please select a session to generate a report."
            )
            return

        # Suggest filename based on session data
        asset = self._current_session_data.get("asset", "unknown").replace("/", "-")
        date_str = QDate.currentDate().toString("yyyyMMdd")
        suggested_name = f"calibration_report_{asset}_{date_str}.pdf"

        file_path, _ = QFileDialog.getSaveFileName(
            self,
            "Save PDF Report",
            suggested_name,
            "PDF Files (*.pdf)",
        )

        if not file_path:
            return

        logger.info(f"Generating PDF: {file_path}")

        try:
            self._generate_pdf_report(file_path)
            QMessageBox.information(
                self,
                "Report Generated",
                f"PDF report saved to:\n{file_path}",
            )
        except Exception as e:
            logger.error(f"Failed to generate PDF: {e}")
            QMessageBox.critical(self, "Error", f"Failed to generate PDF:\n{e}")

    def _generate_pdf_report(self, file_path: str):
        """Generate the actual PDF report using the new template generator."""
        from pathlib import Path
        from calsystem.utils.report_generator import (
            generate_calibration_report,
            build_report_from_session,
        )

        # Build report data from session
        report_data = build_report_from_session(self._current_session_id)

        if not report_data:
            raise Exception("Failed to build report data from session")

        # Generate the PDF
        success = generate_calibration_report(
            output_path=Path(file_path),
            header=report_data['header'],
            sections=report_data['sections'],
        )

        if not success:
            raise Exception("Failed to generate PDF report")

    def _on_export_excel(self):
        """Export to Excel."""
        if not self._current_session_id or not self._current_session_data:
            QMessageBox.warning(
                self, "No Selection", "Please select a session to export."
            )
            return

        asset = self._current_session_data.get("asset", "unknown").replace("/", "-")
        date_str = QDate.currentDate().toString("yyyyMMdd")
        suggested_name = f"calibration_data_{asset}_{date_str}.xlsx"

        file_path, _ = QFileDialog.getSaveFileName(
            self,
            "Export to Excel",
            suggested_name,
            "Excel Files (*.xlsx)",
        )

        if not file_path:
            return

        logger.info(f"Exporting to Excel: {file_path}")

        try:
            self._export_to_excel(file_path)
            QMessageBox.information(
                self,
                "Export Complete",
                f"Data exported to:\n{file_path}",
            )
        except Exception as e:
            logger.error(f"Failed to export to Excel: {e}")
            QMessageBox.critical(self, "Error", f"Failed to export:\n{e}")

    def _export_to_excel(self, file_path: str):
        """Export session data to Excel."""
        try:
            import openpyxl
            from openpyxl.styles import Font, PatternFill, Alignment
        except ImportError:
            raise ImportError("openpyxl is required for Excel export. Install with: pip install openpyxl")

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Calibration Results"

        # Header style
        header_font = Font(bold=True)
        header_fill = PatternFill(start_color="CCCCCC", end_color="CCCCCC", fill_type="solid")

        # Session info
        data = self._current_session_data
        ws['A1'] = "Asset Number:"
        ws['B1'] = data.get("asset", "")
        ws['A2'] = "DUT:"
        ws['B2'] = data.get("dut_info", "")
        ws['A3'] = "Work Order:"
        ws['B3'] = data.get("work_order", "")
        ws['A4'] = "Date:"
        ws['B4'] = data.get("date").strftime("%Y-%m-%d %H:%M") if data.get("date") else ""
        ws['A5'] = "Technician:"
        ws['B5'] = data.get("technician", "")
        ws['A6'] = "Result:"
        ws['B6'] = data.get("result", "")

        for row in range(1, 7):
            ws.cell(row=row, column=1).font = header_font

        # Results header
        start_row = 8
        headers = ["Test Point", "Nominal", "Measured", "Deviation", "Tolerance", "Status"]
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=start_row, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill

        # Results data
        for row in range(self.results_table.rowCount()):
            for col in range(6):
                item = self.results_table.item(row, col)
                ws.cell(row=start_row + row + 1, column=col + 1, value=item.text() if item else "")

        # Auto-size columns
        for col in ws.columns:
            max_length = 0
            column = col[0].column_letter
            for cell in col:
                if cell.value:
                    max_length = max(max_length, len(str(cell.value)))
            ws.column_dimensions[column].width = max_length + 2

        wb.save(file_path)

    def _on_print(self):
        """Print report."""
        selected = self.sessions_table.selectedItems()
        if not selected:
            QMessageBox.warning(
                self, "No Selection", "Please select a session to print."
            )
            return

        logger.info("Printing report")
        # TODO: Implement print functionality
        QMessageBox.information(
            self,
            "Print",
            "Print functionality will be implemented.",
        )
