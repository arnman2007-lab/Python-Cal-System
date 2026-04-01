"""
Excel import wizard dialog for importing procedures from Excel files.
"""

from typing import Optional, List, Dict, Any

from PyQt6.QtWidgets import (
    QDialog,
    QVBoxLayout,
    QHBoxLayout,
    QLabel,
    QPushButton,
    QComboBox,
    QTableWidget,
    QTableWidgetItem,
    QHeaderView,
    QStackedWidget,
    QWidget,
    QGroupBox,
    QMessageBox,
    QCheckBox,
)
from PyQt6.QtCore import Qt
from loguru import logger

try:
    import openpyxl
    OPENPYXL_AVAILABLE = True
except ImportError:
    OPENPYXL_AVAILABLE = False


class ExcelImportDialog(QDialog):
    """Wizard dialog for importing calibration procedures from Excel files."""

    def __init__(self, file_path: str, parent=None):
        super().__init__(parent)
        self.file_path = file_path
        self.workbook = None
        self.sheet_data: List[List[Any]] = []
        self.column_mappings: Dict[str, int] = {}
        self.imported_rows: List[Dict[str, Any]] = []

        self.setWindowTitle("Import Procedure from Excel")
        self.setMinimumSize(700, 500)
        self._init_ui()
        self._load_workbook()

    def _init_ui(self):
        """Initialize the UI."""
        layout = QVBoxLayout(self)

        # Stacked widget for wizard steps
        self.stack = QStackedWidget()
        layout.addWidget(self.stack)

        # Step 1: Select sheet
        self._create_sheet_step()

        # Step 2: Map columns
        self._create_mapping_step()

        # Step 3: Preview
        self._create_preview_step()

        # Navigation buttons
        nav_layout = QHBoxLayout()
        self.back_btn = QPushButton("< Back")
        self.back_btn.clicked.connect(self._on_back)
        self.back_btn.setEnabled(False)
        nav_layout.addWidget(self.back_btn)

        nav_layout.addStretch()

        self.next_btn = QPushButton("Next >")
        self.next_btn.clicked.connect(self._on_next)
        nav_layout.addWidget(self.next_btn)

        self.import_btn = QPushButton("Import")
        self.import_btn.clicked.connect(self._on_import)
        self.import_btn.hide()
        nav_layout.addWidget(self.import_btn)

        self.cancel_btn = QPushButton("Cancel")
        self.cancel_btn.clicked.connect(self.reject)
        nav_layout.addWidget(self.cancel_btn)

        layout.addLayout(nav_layout)

    def _create_sheet_step(self):
        """Create step 1: sheet selection."""
        page = QWidget()
        layout = QVBoxLayout(page)

        layout.addWidget(QLabel("<h3>Step 1: Select Sheet</h3>"))
        layout.addWidget(QLabel(f"File: {self.file_path}"))

        self.sheet_combo = QComboBox()
        self.sheet_combo.currentIndexChanged.connect(self._on_sheet_changed)
        layout.addWidget(QLabel("Select worksheet:"))
        layout.addWidget(self.sheet_combo)

        # Preview table
        layout.addWidget(QLabel("Preview:"))
        self.sheet_preview_table = QTableWidget()
        self.sheet_preview_table.setMaximumHeight(200)
        layout.addWidget(self.sheet_preview_table)

        layout.addStretch()
        self.stack.addWidget(page)

    def _create_mapping_step(self):
        """Create step 2: column mapping."""
        page = QWidget()
        layout = QVBoxLayout(page)

        layout.addWidget(QLabel("<h3>Step 2: Map Columns</h3>"))
        layout.addWidget(QLabel("Map Excel columns to test point fields:"))

        # Mapping group
        mapping_group = QGroupBox("Column Mappings")
        mapping_layout = QVBoxLayout(mapping_group)

        # Create mapping combos for each field
        self.mapping_combos: Dict[str, QComboBox] = {}
        fields = [
            ("section", "Section Name"),
            ("description", "Test Point Description"),
            ("nominal", "Nominal Value"),
            ("unit", "Unit"),
            ("tolerance", "Tolerance"),
            ("tolerance_type", "Tolerance Type (%, Abs, PPM)"),
            ("frequency", "Frequency (optional)"),
        ]

        for field_key, field_label in fields:
            row = QHBoxLayout()
            row.addWidget(QLabel(f"{field_label}:"))
            combo = QComboBox()
            combo.addItem("-- Skip --", -1)
            self.mapping_combos[field_key] = combo
            row.addWidget(combo)
            mapping_layout.addLayout(row)

        layout.addWidget(mapping_group)

        # Options
        self.header_check = QCheckBox("First row is header (skip it)")
        self.header_check.setChecked(True)
        layout.addWidget(self.header_check)

        layout.addStretch()
        self.stack.addWidget(page)

    def _create_preview_step(self):
        """Create step 3: preview."""
        page = QWidget()
        layout = QVBoxLayout(page)

        layout.addWidget(QLabel("<h3>Step 3: Preview Import</h3>"))
        layout.addWidget(QLabel("Review the data that will be imported:"))

        self.preview_table = QTableWidget()
        self.preview_table.setColumnCount(6)
        self.preview_table.setHorizontalHeaderLabels([
            "Section", "Description", "Nominal", "Unit", "Tolerance", "Tol Type"
        ])
        self.preview_table.horizontalHeader().setSectionResizeMode(QHeaderView.ResizeMode.Stretch)
        layout.addWidget(self.preview_table)

        self.preview_summary = QLabel("")
        layout.addWidget(self.preview_summary)

        layout.addStretch()
        self.stack.addWidget(page)

    def _load_workbook(self):
        """Load the Excel workbook."""
        if not OPENPYXL_AVAILABLE:
            QMessageBox.critical(
                self,
                "Missing Dependency",
                "openpyxl is required for Excel import.\n\nInstall it with: pip install openpyxl"
            )
            self.reject()
            return

        try:
            self.workbook = openpyxl.load_workbook(self.file_path, read_only=True, data_only=True)
            for sheet_name in self.workbook.sheetnames:
                self.sheet_combo.addItem(sheet_name)

            if self.workbook.sheetnames:
                self._on_sheet_changed(0)

        except Exception as e:
            logger.error(f"Failed to load workbook: {e}")
            QMessageBox.critical(self, "Error", f"Failed to load Excel file:\n{e}")
            self.reject()

    def _on_sheet_changed(self, index: int):
        """Handle sheet selection change."""
        if not self.workbook or index < 0:
            return

        sheet_name = self.sheet_combo.currentText()
        sheet = self.workbook[sheet_name]

        # Load sheet data
        self.sheet_data = []
        for row in sheet.iter_rows(max_row=100, values_only=True):
            self.sheet_data.append(list(row))

        # Update preview
        self._update_sheet_preview()

        # Update column mappings
        self._update_column_options()

    def _update_sheet_preview(self):
        """Update sheet preview table."""
        self.sheet_preview_table.clear()

        if not self.sheet_data:
            return

        # Show first 10 rows
        preview_rows = min(10, len(self.sheet_data))
        num_cols = max(len(row) for row in self.sheet_data[:preview_rows]) if self.sheet_data else 0

        self.sheet_preview_table.setRowCount(preview_rows)
        self.sheet_preview_table.setColumnCount(num_cols)

        # Set column headers as A, B, C, etc.
        headers = [chr(65 + i) if i < 26 else f"Col{i+1}" for i in range(num_cols)]
        self.sheet_preview_table.setHorizontalHeaderLabels(headers)

        for i, row in enumerate(self.sheet_data[:preview_rows]):
            for j, value in enumerate(row):
                if j < num_cols:
                    item = QTableWidgetItem(str(value) if value is not None else "")
                    self.sheet_preview_table.setItem(i, j, item)

    def _update_column_options(self):
        """Update column options in mapping combos."""
        if not self.sheet_data:
            return

        # Get column count from first row
        num_cols = len(self.sheet_data[0]) if self.sheet_data else 0

        # Get header names from first row
        headers = []
        if self.sheet_data:
            for i, val in enumerate(self.sheet_data[0]):
                col_letter = chr(65 + i) if i < 26 else f"Col{i+1}"
                header_text = str(val) if val else col_letter
                headers.append(f"{col_letter}: {header_text[:20]}")

        for combo in self.mapping_combos.values():
            current = combo.currentData()
            combo.clear()
            combo.addItem("-- Skip --", -1)
            for i, header in enumerate(headers):
                combo.addItem(header, i)

            # Restore selection if valid
            if current is not None and current >= 0 and current < len(headers):
                combo.setCurrentIndex(current + 1)

    def _on_back(self):
        """Go to previous step."""
        current = self.stack.currentIndex()
        if current > 0:
            self.stack.setCurrentIndex(current - 1)
            self._update_nav_buttons()

    def _on_next(self):
        """Go to next step."""
        current = self.stack.currentIndex()

        if current == 1:
            # Validate mappings before preview
            if not self._validate_mappings():
                return
            self._generate_preview()

        if current < self.stack.count() - 1:
            self.stack.setCurrentIndex(current + 1)
            self._update_nav_buttons()

    def _update_nav_buttons(self):
        """Update navigation button states."""
        current = self.stack.currentIndex()
        self.back_btn.setEnabled(current > 0)
        self.next_btn.setVisible(current < self.stack.count() - 1)
        self.import_btn.setVisible(current == self.stack.count() - 1)

    def _validate_mappings(self) -> bool:
        """Validate that required mappings are set."""
        nominal_idx = self.mapping_combos["nominal"].currentData()

        if nominal_idx is None or nominal_idx < 0:
            QMessageBox.warning(
                self,
                "Missing Mapping",
                "Please map at least the Nominal Value column."
            )
            return False

        return True

    def _generate_preview(self):
        """Generate preview data from mappings."""
        self.imported_rows = []
        self.preview_table.setRowCount(0)

        if not self.sheet_data:
            return

        # Get column indices
        section_idx = self.mapping_combos["section"].currentData()
        desc_idx = self.mapping_combos["description"].currentData()
        nominal_idx = self.mapping_combos["nominal"].currentData()
        unit_idx = self.mapping_combos["unit"].currentData()
        tol_idx = self.mapping_combos["tolerance"].currentData()
        tol_type_idx = self.mapping_combos["tolerance_type"].currentData()

        # Start row (skip header if checked)
        start_row = 1 if self.header_check.isChecked() else 0

        for row in self.sheet_data[start_row:]:
            # Extract values with bounds checking
            def get_val(idx):
                if idx is None or idx < 0 or idx >= len(row):
                    return None
                return row[idx]

            nominal = get_val(nominal_idx)
            if nominal is None:
                continue  # Skip rows without nominal value

            try:
                nominal_float = float(nominal)
            except (ValueError, TypeError):
                continue  # Skip non-numeric

            row_data = {
                "section": str(get_val(section_idx) or "Default"),
                "description": str(get_val(desc_idx) or ""),
                "nominal": nominal_float,
                "unit": str(get_val(unit_idx) or "V"),
                "tolerance": float(get_val(tol_idx) or 0),
                "tolerance_type": str(get_val(tol_type_idx) or "percent"),
            }
            self.imported_rows.append(row_data)

            # Add to preview table
            row_idx = self.preview_table.rowCount()
            self.preview_table.insertRow(row_idx)
            self.preview_table.setItem(row_idx, 0, QTableWidgetItem(row_data["section"]))
            self.preview_table.setItem(row_idx, 1, QTableWidgetItem(row_data["description"]))
            self.preview_table.setItem(row_idx, 2, QTableWidgetItem(str(row_data["nominal"])))
            self.preview_table.setItem(row_idx, 3, QTableWidgetItem(row_data["unit"]))
            self.preview_table.setItem(row_idx, 4, QTableWidgetItem(str(row_data["tolerance"])))
            self.preview_table.setItem(row_idx, 5, QTableWidgetItem(row_data["tolerance_type"]))

        self.preview_summary.setText(f"Found {len(self.imported_rows)} test points to import.")

    def _on_import(self):
        """Perform the import."""
        if not self.imported_rows:
            QMessageBox.warning(self, "No Data", "No test points to import.")
            return

        self.accept()

    def get_imported_data(self) -> List[Dict[str, Any]]:
        """Get the imported data."""
        return self.imported_rows
