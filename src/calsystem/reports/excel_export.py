"""
Excel export module using openpyxl.
"""

from typing import Optional, List, Dict, Any
from pathlib import Path
from datetime import datetime

from loguru import logger

try:
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
    from openpyxl.utils import get_column_letter
    OPENPYXL_AVAILABLE = True
except ImportError:
    OPENPYXL_AVAILABLE = False
    logger.warning("openpyxl not available - Excel export disabled")


class ExcelExporter:
    """Exports calibration data to Excel format."""

    def __init__(self):
        self._thin_border = None
        if OPENPYXL_AVAILABLE:
            self._thin_border = Border(
                left=Side(style="thin"),
                right=Side(style="thin"),
                top=Side(style="thin"),
                bottom=Side(style="thin"),
            )

    def export(
        self,
        output_path: str,
        session_data: Dict[str, Any],
        test_results: List[Dict[str, Any]],
        template_path: Optional[str] = None,
    ) -> bool:
        """
        Export calibration data to Excel.

        Args:
            output_path: Path to save the Excel file.
            session_data: Calibration session information.
            test_results: List of test point results.
            template_path: Optional template workbook to populate.

        Returns:
            True if successful.
        """
        if not OPENPYXL_AVAILABLE:
            logger.error("openpyxl not available")
            return False

        try:
            if template_path and Path(template_path).exists():
                return self._export_to_template(
                    output_path, session_data, test_results, template_path
                )
            else:
                return self._export_standard(output_path, session_data, test_results)

        except Exception as e:
            logger.error(f"Excel export error: {e}")
            return False

    def _export_standard(
        self,
        output_path: str,
        session_data: Dict[str, Any],
        test_results: List[Dict[str, Any]],
    ) -> bool:
        """Export to standard format (new workbook)."""
        wb = Workbook()
        ws = wb.active
        ws.title = "Calibration Data"

        # Header styles
        header_font = Font(bold=True, size=12)
        header_fill = PatternFill(start_color="CCCCCC", end_color="CCCCCC", fill_type="solid")

        # Title
        ws["A1"] = "Calibration Report"
        ws["A1"].font = Font(bold=True, size=16)
        ws.merge_cells("A1:F1")

        # Session info
        row = 3
        info_items = [
            ("Asset Number:", session_data.get("asset_number", "")),
            ("Device:", f"{session_data.get('make', '')} {session_data.get('model', '')}"),
            ("Serial Number:", session_data.get("serial_number", "")),
            ("Work Order:", session_data.get("work_order", "")),
            ("Date:", session_data.get("date", datetime.now().strftime("%Y-%m-%d"))),
            ("Technician:", session_data.get("technician", "")),
        ]

        for label, value in info_items:
            ws.cell(row=row, column=1, value=label).font = Font(bold=True)
            ws.cell(row=row, column=2, value=value)
            row += 1

        # Results table header
        row += 1
        headers = ["Test Point", "Nominal", "Measured", "Deviation", "Tolerance", "Status"]
        for col, header in enumerate(headers, start=1):
            cell = ws.cell(row=row, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.border = self._thin_border
            cell.alignment = Alignment(horizontal="center")

        # Results data
        row += 1
        pass_fill = PatternFill(start_color="90EE90", end_color="90EE90", fill_type="solid")
        fail_fill = PatternFill(start_color="FFB6C1", end_color="FFB6C1", fill_type="solid")

        for result in test_results:
            ws.cell(row=row, column=1, value=result.get("test_point", ""))
            ws.cell(row=row, column=2, value=result.get("nominal", ""))
            ws.cell(row=row, column=3, value=result.get("measured", ""))
            ws.cell(row=row, column=4, value=result.get("deviation", ""))
            ws.cell(row=row, column=5, value=result.get("tolerance", ""))

            status_cell = ws.cell(row=row, column=6, value=result.get("status", ""))
            status = result.get("status", "").upper()
            if status == "PASS":
                status_cell.fill = pass_fill
            elif status == "FAIL":
                status_cell.fill = fail_fill

            # Apply borders
            for col in range(1, 7):
                ws.cell(row=row, column=col).border = self._thin_border

            row += 1

        # Adjust column widths
        column_widths = [20, 15, 15, 12, 12, 10]
        for i, width in enumerate(column_widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = width

        # Summary
        row += 2
        total = len(test_results)
        passed = sum(1 for r in test_results if r.get("status", "").upper() == "PASS")

        ws.cell(row=row, column=1, value="Summary:").font = Font(bold=True)
        row += 1
        ws.cell(row=row, column=1, value=f"Total: {total}, Passed: {passed}, Failed: {total - passed}")

        # Save
        wb.save(output_path)
        logger.info(f"Excel export saved: {output_path}")
        return True

    def _export_to_template(
        self,
        output_path: str,
        session_data: Dict[str, Any],
        test_results: List[Dict[str, Any]],
        template_path: str,
    ) -> bool:
        """Export to existing template with cell mappings."""
        wb = load_workbook(template_path)

        # Get mappings from test results (each result may have excel_sheet and excel_cell)
        for result in test_results:
            sheet_name = result.get("excel_sheet")
            cell_address = result.get("excel_cell")
            measured_value = result.get("measured")

            if sheet_name and cell_address and measured_value:
                if sheet_name in wb.sheetnames:
                    ws = wb[sheet_name]
                    ws[cell_address] = measured_value
                    logger.debug(f"Wrote {measured_value} to {sheet_name}!{cell_address}")

        # Save as new file (don't overwrite template)
        wb.save(output_path)
        logger.info(f"Template export saved: {output_path}")
        return True

    def export_with_mappings(
        self,
        output_path: str,
        mappings: List[Dict[str, Any]],
        template_path: Optional[str] = None,
    ) -> bool:
        """
        Export data using specific cell mappings.

        Args:
            output_path: Path to save.
            mappings: List of {sheet, cell, value} dictionaries.
            template_path: Optional template to start from.

        Returns:
            True if successful.
        """
        if not OPENPYXL_AVAILABLE:
            return False

        try:
            if template_path and Path(template_path).exists():
                wb = load_workbook(template_path)
            else:
                wb = Workbook()

            for mapping in mappings:
                sheet_name = mapping.get("sheet", "Sheet1")
                cell_address = mapping.get("cell")
                value = mapping.get("value")

                if not cell_address:
                    continue

                # Create sheet if needed
                if sheet_name not in wb.sheetnames:
                    wb.create_sheet(sheet_name)

                ws = wb[sheet_name]
                ws[cell_address] = value

            wb.save(output_path)
            logger.info(f"Mapped export saved: {output_path}")
            return True

        except Exception as e:
            logger.error(f"Mapped export error: {e}")
            return False
